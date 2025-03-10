import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import io
import os
import re
from typing import List, Dict, Any, Optional

def create_business_definition_excel(all_tables: List[Dict[str, Any]], pdf_filename: str) -> Optional[io.BytesIO]:
    """
    테이블 데이터로 업무정의서 형식의 엑셀 생성

    Args:
        all_tables: 추출된 테이블 정보 리스트
        pdf_filename: PDF 파일명
        
    Returns:
        io.BytesIO: 메모리에 저장된 Excel 파일
    """
    if not all_tables:
        return None
    
    # 파일명에서 상품명 추출
    product_name = os.path.splitext(os.path.basename(pdf_filename))[0]
    
    # Workbook 생성
    wb = Workbook()
    ws = wb.active
    ws.title = "보장내용 개정사항"
    
    # 스타일 정의
    header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")  # 노란 형광색용
    gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")    # 회색용
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 제목 행 추가
    ws.merge_cells('A1:F1')
    cell = ws.cell(row=1, column=1, value=f"{product_name} - 보장내용 개정사항")
    cell.font = Font(size=14, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 헤더 행 추가
    headers = ["페이지", "보장명", "지급사유", "지급금액", "비고", "강조여부"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 테이블 데이터 추가
    row_idx = 4
    for table_info in all_tables:
        df = table_info['df']
        page_num = table_info.get('page', 0)
        
        # 강조색 및 복원된 텍스트 정보
        highlight_cells = table_info.get('highlight_cells', [])
        missing_text_cells = table_info.get('missing_text_cells', [])
        gray_cells = table_info.get('gray_cells', [])
        
        # 데이터프레임 행, 열 수 확인
        rows, cols = df.shape
        
        # 모든 행 포함 (모든 행을 이제 포함하여 형광펜만 있는 행도 표시)
        for df_row_idx in range(rows):
            # 해당 행이 강조 처리된 행인지 확인
            row_has_highlight = any((df_row_idx, c) in highlight_cells for c in range(cols))
            
            # 셀 값 준비
            row_values = [page_num]  # 첫 번째 열은 페이지 번호
            
            # 데이터프레임의 해당 행 값 추가 (컬럼 최대 4개까지만)
            for df_col_idx in range(min(cols, 4)):
                cell_value = df.iloc[df_row_idx, df_col_idx]
                if pd.isna(cell_value):
                    cell_value = ""
                row_values.append(str(cell_value))
            
            # 부족한 열은 빈 문자열로 채움
            while len(row_values) < 5:
                row_values.append("")
            
            # 강조여부 열 추가
            highlight_status = []
            if row_has_highlight:
                highlight_status.append("강조")
            if any((df_row_idx, c) in gray_cells for c in range(cols)):
                highlight_status.append("취소선")
            if any((df_row_idx, c) in missing_text_cells for c in range(cols)):
                highlight_status.append("형광펜 텍스트 복원")
            
            row_values.append(", ".join(highlight_status) if highlight_status else "")
            
            # 엑셀에 행 추가
            for col_idx, value in enumerate(row_values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                
                # 강조색 적용 (열이 4개 이상인 경우에도 정확히 반영)
                df_col = col_idx - 2 if col_idx >= 2 and col_idx <= min(cols+1, 5) else -1
                
                # 데이터 셀에만 형식 적용
                if 2 <= col_idx <= 5 and df_col >= 0:
                    # 강조색 적용
                    if (df_row_idx, df_col) in highlight_cells:
                        cell.fill = yellow_fill
                    # 회색 적용 (취소선 있는 경우)
                    elif (df_row_idx, df_col) in gray_cells:
                        cell.fill = gray_fill
                        # 취소선 텍스트 스타일 적용
                        cell.font = Font(strike=True)
            
            row_idx += 1
    
    # 열 너비 조정
    ws.column_dimensions['A'].width = 10  # 페이지
    ws.column_dimensions['B'].width = 25  # 보장명
    ws.column_dimensions['C'].width = 40  # 지급사유
    ws.column_dimensions['D'].width = 25  # 지급금액
    ws.column_dimensions['E'].width = 15  # 비고
    ws.column_dimensions['F'].width = 15  # 강조여부
    
    # 엑셀 파일 저장 (메모리에)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output

def map_highlights_to_tables(tables: List[Dict[str, Any]], highlight_cells_mapping: Dict[int, List[tuple]]) -> List[Dict[str, Any]]:
    """
    테이블 셀에 강조색 정보 매핑
    
    Args:
        tables: 테이블 리스트
        highlight_cells_mapping: 페이지별 강조색 셀 좌표 매핑
        
    Returns:
        List[Dict[str, Any]]: 강조색 정보가 매핑된 테이블 리스트
    """
    for table_info in tables:
        page_num = table_info.get('page', 0)
        
        # 해당 페이지의 강조색 셀 좌표가 있는 경우
        if page_num in highlight_cells_mapping:
            highlight_cells = []
            gray_cells = []
            
            # 강조색/회색 셀 정보 저장
            for cell_type, coords in highlight_cells_mapping[page_num]:
                if cell_type == 'highlight':
                    highlight_cells.append(coords)
                elif cell_type == 'gray':
                    gray_cells.append(coords)
            
            # 테이블 정보에 강조 셀 정보 추가
            table_info['highlight_cells'] = highlight_cells
            table_info['gray_cells'] = gray_cells
    
    return tables