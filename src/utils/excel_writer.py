from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import MergedCell
import logging
from pathlib import Path
import pandas as pd
from typing import List, Optional, Dict

class ExcelWriter:
    def __init__(self, output_path: str):
        self.output_path = Path(output_path)
        self.wb = Workbook()
        self.ws = self.wb.active
        self.current_row = 1
        self._init_styles()
        self.logger = logging.getLogger(__name__)

    def _init_styles(self):
        # 기본 스타일
        self.default_fill = PatternFill(fill_type=None)
        
        # 헤더 스타일
        self.header_fill = PatternFill(start_color='E6E6E6', end_color='E6E6E6', fill_type='solid')
        self.header_font = Font(bold=True)
        
        # 강조 스타일
        self.highlight_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
        
        # 변경 유형별 스타일
        self.highlight_fills = {
            'default': PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid'),  # 노란색
            'added': PatternFill(start_color='E3FFE3', end_color='E3FFE3', fill_type='solid'),    # 연한 녹색
            'deleted': PatternFill(start_color='FFE3E3', end_color='FFE3E3', fill_type='solid'),  # 연한 빨간색
            'modified': PatternFill(start_color='FFE3B3', end_color='FFE3B3', fill_type='solid')  # 연한 주황색
        }
        
        # 테두리
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 정렬
        self.alignment = Alignment(wrap_text=True, vertical='center')

    def write_table(self, df: pd.DataFrame, sheet_name: Optional[str] = None, 
                highlights: Optional[List[List[bool]]] = None,
                change_types: Optional[List[List[str]]] = None,  # 추가
                metadata: Optional[Dict] = None):
        try:
            # 시트 설정
            if sheet_name:
                if sheet_name in self.wb.sheetnames:
                    self.ws = self.wb[sheet_name]
                else:
                    self.ws = self.wb.create_sheet(sheet_name)

            # [변경] 태그 제거
            df = df.apply(lambda x: x.str.replace(r'\s*\[변경\]', '') 
                        if pd.api.types.is_string_dtype(x) else x)

            # 메타데이터 작성
            if metadata:
                for key, value in metadata.items():
                    cell = self.ws.cell(row=self.current_row, column=1, value=f"{key}: {value}")
                    cell.font = Font(italic=True)
                    self.current_row += 1
                self.current_row += 1

            # 헤더 작성
            for col_idx, col_name in enumerate(df.columns, 1):
                cell = self.ws.cell(row=self.current_row, column=col_idx, value=col_name)
                cell.fill = self.header_fill
                cell.font = self.header_font
                cell.border = self.border
                cell.alignment = self.alignment

            self.current_row += 1

            # 데이터 작성
            for row_idx, row in enumerate(df.itertuples(index=False), 0):
                excel_row = self.current_row + row_idx
                for col_idx, value in enumerate(row, 1):
                    cell = self.ws.cell(row=excel_row, column=col_idx, value=value)
                    cell.border = self.border
                    cell.alignment = self.alignment
                    
                    # A열은 색상 없음
                    if col_idx == 1:
                        cell.fill = self.default_fill
                    # 하이라이트 적용
                    elif highlights and row_idx < len(highlights) and col_idx-1 < len(highlights[row_idx]):
                        if highlights[row_idx][col_idx-1]:
                            if change_types and row_idx < len(change_types) and col_idx-1 < len(change_types[row_idx]):
                                change_type = change_types[row_idx][col_idx-1]
                                cell.fill = self.highlight_fills.get(change_type, self.highlight_fill)
                            else:
                                cell.fill = self.highlight_fill

            self.current_row += len(df)
            self._adjust_column_widths()
            
        except Exception as e:
            self.logger.error(f"Failed to write table: {e}")
            raise

    def _adjust_column_widths(self):
        for column_cells in self.ws.columns:
            max_length = 0
            column_letter = None
            
            # 첫 번째 비병합 셀에서 column_letter 가져오기
            for cell in column_cells:
                if not isinstance(cell, MergedCell):
                    column_letter = cell.column_letter
                    break
            
            if column_letter is None:
                continue
                
            # 컬럼 내 모든 셀의 최대 길이 계산
            for cell in column_cells:
                try:
                    if cell.value and not isinstance(cell, MergedCell):
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
                    
            adjusted_width = min(max_length + 2, 50)
            self.ws.column_dimensions[column_letter].width = adjusted_width

    def write_section_header(self, title: str, span: int = 5):
        cell = self.ws.cell(row=self.current_row, column=1, value=title)
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill(start_color='E6E6E6', end_color='E6E6E6', fill_type='solid')
        
        if span > 1:
            merge_range = f"A{self.current_row}:{get_column_letter(span)}{self.current_row}"
            self.ws.merge_cells(merge_range)
        
        self.current_row += 2

    def save(self):
        try:
            self.output_path.parent.mkdir(parents=True, exist_ok=True)
            self.wb.save(self.output_path)
            self.logger.info(f"Saved Excel file: {self.output_path}")
        except Exception as e:
            self.logger.error(f"Failed to save file: {e}")
            raise